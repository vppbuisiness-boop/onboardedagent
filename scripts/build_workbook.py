"""Build the Edgeline model workbook: board, bettable picks, slips, real-line record, backtests, graded lines, model reference.

Usage: python scripts/build_workbook.py [out.xlsx]   (default data/exports/edgeline_model.xlsx)
"""
from __future__ import annotations

import datetime as dt
import glob
import json
import os
import sqlite3
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from edgeline.config import DEFAULT_MAX_LINE_MOVE, DEFAULT_MIN_EV, DEFAULT_MIN_PROB, DEFAULT_MARKET_SHRINK, UNPROVEN_MARKETS  # noqa: E402
from edgeline.ev.payouts import LADDERS  # noqa: E402
from edgeline.features.build import FEATURE_COLUMNS  # noqa: E402
from edgeline.grading.results import results_frame  # noqa: E402
from edgeline.grading.roi import parlay_roi, wilson  # noqa: E402

DB = Path("data/edgeline.db")
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("data/exports/edgeline_model.xlsx")

NAVY = "1F3864"; TEAL = "0F766E"; GREEN = "C6EFCE"; RED = "FFC7CE"; AMBER = "FFEB9C"; GREY = "F2F2F2"; BLUE = "DDEBF7"; ORANGE = "FCE4D6"
SPORT_FILL = {"dota": "E2EFDA", "lol": "DDEBF7", "cs2": "FFF2CC", "val": "FCE4D6", "cod": "EDEDED"}
HEAD_FONT = Font(bold=True, color="FFFFFF"); BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF"); BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def header(ws, cols, row=1, fill=NAVY):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = HEAD_FONT; cell.fill = PatternFill("solid", fgColor=fill); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = BORDER
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    ws.row_dimensions[row].height = 30


def write_df(ws, df, start_row=1, fill=NAVY, pct_cols=(), widths=None, number_formats=None):
    header(ws, list(df.columns), start_row, fill)
    for i, rec in enumerate(df.itertuples(index=False), start_row + 1):
        for j, v in enumerate(rec, 1):
            if isinstance(v, float) and pd.isna(v):
                v = None
            cell = ws.cell(row=i, column=j, value=v); cell.border = BORDER
            name = df.columns[j - 1]
            if name in pct_cols:
                cell.number_format = "0.0%"
            elif number_formats and name in number_formats:
                cell.number_format = number_formats[name]
    for j, name in enumerate(df.columns, 1):
        w = (widths or {}).get(name) or min(max(len(str(name)), *(len(str(x)) for x in df[name].head(200).tolist() or [""])) + 2, 48)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(df.columns))}{start_row + len(df)}"
    return start_row + len(df)


def color_sport(ws, df, start_row, col_name="sport"):
    if col_name not in df.columns:
        return
    j = list(df.columns).index(col_name) + 1
    for i, sp in enumerate(df[col_name].tolist(), start_row + 1):
        f = SPORT_FILL.get(str(sp).lower())
        if f:
            ws.cell(row=i, column=j).fill = PatternFill("solid", fgColor=f)


def to_et(s):
    try:
        return pd.to_datetime(s, utc=True).tz_convert("America/New_York").strftime("%a %m-%d %I:%M %p ET")
    except Exception:
        return s


def main():
    conn = sqlite3.connect(DB, timeout=60)
    wb = Workbook()
    now = dt.datetime.now(dt.timezone.utc)

    # ---------------- Board ----------------
    board = pd.read_sql_query("""
        SELECT l.sport, l.start_time, l.player_name AS player, l.team, l.opponent, l.stat_type, l.current_line AS line, l.open_line,
               p.projection, p.lean, p.prob, p.ev, p.bettable, l.league, p.notes, l.projection_id
        FROM lines l JOIN predictions p ON p.book=l.book AND p.projection_id=l.projection_id
        WHERE l.book='prizepicks' AND l.start_time > strftime('%Y-%m-%dT%H:%M:%SZ','now')
        ORDER BY p.bettable DESC, p.prob DESC""", conn)
    board["start (ET)"] = board["start_time"].map(to_et)
    board = board[["sport", "start (ET)", "player", "team", "opponent", "stat_type", "line", "open_line", "projection", "lean", "prob", "ev", "bettable", "notes", "projection_id"]]
    board["bettable"] = board["bettable"].map({1: "YES", 0: "no"})
    ws = wb.active; ws.title = "Board"
    ws["A1"] = f"Edgeline board: every upcoming PrizePicks line priced by the model (as of {now:%Y-%m-%d %H:%M} UTC). Green rows = bettable (>= {DEFAULT_MIN_PROB:.0%} probability and >= {DEFAULT_MIN_EV:.0%} EV, standard odds, not voidable, not moved > {DEFAULT_MAX_LINE_MOVE:.0%} against)."
    ws["A1"].font = Font(bold=True, size=12, color=NAVY)
    last = write_df(ws, board, start_row=3, pct_cols=("prob", "ev"), number_formats={"projection": "0.0", "line": "0.0", "open_line": "0.0"})
    color_sport(ws, board, 3)
    n = len(board)
    if n:
        rng = f"A4:{get_column_letter(len(board.columns))}{3 + n}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=['$M4="YES"'], fill=PatternFill("solid", fgColor=GREEN)))
        ws.conditional_formatting.add(f"K4:K{3 + n}", ColorScaleRule(start_type="num", start_value=0.5, start_color="FFFFFF", end_type="num", end_value=0.8, end_color="63BE7B"))
        ws.conditional_formatting.add(f"L4:L{3 + n}", ColorScaleRule(start_type="num", start_value=-0.1, start_color="F8696B", mid_type="num", mid_value=0.05, mid_color="FFFFFF", end_type="num", end_value=0.4, end_color="63BE7B"))
        ws.conditional_formatting.add(f"J4:J{3 + n}", CellIsRule(operator="equal", formula=['"OVER"'], fill=PatternFill("solid", fgColor=BLUE)))
        ws.conditional_formatting.add(f"J4:J{3 + n}", CellIsRule(operator="equal", formula=['"UNDER"'], fill=PatternFill("solid", fgColor=ORANGE)))

    # ---------------- Bettable ----------------
    bet = board[board["bettable"] == "YES"].drop(columns=["bettable", "projection_id"]).sort_values("prob", ascending=False)
    ws = wb.create_sheet("Bettable picks")
    ws["A1"] = "Bettable lines only, ranked by model probability. Evidence order by market: Dota, Valorant, CS2 headshots, CS2 kills, LoL, COD (see Backtest and Record sheets)."
    ws["A1"].font = Font(bold=True, size=12, color=TEAL)
    write_df(ws, bet, start_row=3, fill=TEAL, pct_cols=("prob", "ev"), number_formats={"projection": "0.0", "line": "0.0", "open_line": "0.0"})
    color_sport(ws, bet, 3)
    if len(bet):
        ws.conditional_formatting.add(f"K4:K{3 + len(bet)}", ColorScaleRule(start_type="num", start_value=0.6, start_color="FFFFFF", end_type="num", end_value=0.8, end_color="63BE7B"))

    # ---------------- Slips ----------------
    try:
        from edgeline.slips.builder import build
        rows = []
        for size in (2, 3, 4):
            for i, s in enumerate(build(conn, "prizepicks", "POWER", size, max_slips=3), 1):
                for k, leg in enumerate(s.legs, 1):
                    rows.append({"slip": f"{size}-pick POWER #{i}", "leg": k, "sport": leg.get("sport"), "player": leg.get("player"), "stat_type": leg.get("stat_type"),
                                 "lean": leg.get("lean"), "line": leg.get("line"), "leg prob": leg.get("prob"), "leg ev": leg.get("ev"), "projection": leg.get("projection"),
                                 "matchup": f"{leg.get('team')} vs {leg.get('opponent')}", "slip hit prob": s.hit_prob, "slip EV": s.ev, "link": s.link})
        slips = pd.DataFrame(rows)
    except Exception as exc:  # the builder needs models and lines; never fail the workbook
        slips = pd.DataFrame([{"slip": f"builder unavailable: {exc}"}])
    ws = wb.create_sheet("Slips")
    ws["A1"] = "Slips the builder proposes from the bettable list (one leg per game, joint probability from the copula). 2-pick pays 3x, 3-pick 5x (6x total return shown as EV), 4-pick 10x."
    ws["A1"].font = Font(bold=True, size=12, color=TEAL)
    write_df(ws, slips, start_row=3, fill=TEAL, pct_cols=("leg prob", "leg ev", "slip hit prob", "slip EV"), number_formats={"projection": "0.0", "line": "0.0"})
    color_sport(ws, slips, 3)

    # ---------------- Line shopping ----------------
    try:
        from edgeline.books.shop import shop
        sh = shop(conn)
    except Exception as exc:
        sh = pd.DataFrame([{"note": f"unavailable: {exc}"}])
    if sh.empty:
        sh = pd.DataFrame([{"note": "no player posted by more than one book right now"}])
    ws = wb.create_sheet("Line shopping")
    ws["A1"] = "Same player and stat on more than one book: each book's line, the model's lean/probability/EV per book, and the best book for the model's lean."
    ws["A1"].font = Font(bold=True, size=12, color=TEAL)
    pct = tuple(c for c in sh.columns if c.endswith(" prob") or c.endswith(" EV") or c == "best_EV")
    write_df(ws, sh, start_row=3, fill=TEAL, pct_cols=pct)
    color_sport(ws, sh, 3)

    # ---------------- Record ----------------
    rf = results_frame(conn, "prizepicks")
    rec_rows = []
    if not rf.empty:
        d = rf.dropna(subset=["win_open"]).copy()
        d["stat"] = d["stat_type"].str.lower().str.replace(r"^map[s]? ?[0-9-]+ ", "", regex=True)
        d["maps"] = d["stat_type"].str.extract(r"MAPS? (\d+(?:-\d+)?)")[0]
        def add(name, g):
            n, w = len(g), int(g["win_open"].sum())
            if not n:
                return
            p, lo, hi = wilson(w, n)
            rec_rows.append({"slice": name, "settled": n, "wins": w, "losses": n - w, "hit rate": p, "CI low": lo, "CI high": hi,
                             "4-pick ROI": parlay_roi(p), "4-pick ROI CI low": parlay_roi(lo), "above break-even (56.2%)": "YES" if lo > 0.5623 else "no", "above 60% target": "YES" if lo > 0.60 else "no"})
        add("ALL model leans", d); add("ALL bettable picks", d[d["bettable"] == 1])
        for sp, g in d.groupby("sport"):
            add(f"{sp.upper()} leans", g); add(f"{sp.upper()} bettable", g[g["bettable"] == 1])
            for st, gg in g.groupby("stat"):
                add(f"{sp.upper()} {st} leans", gg)
            for lean, gg in g.groupby("lean"):
                add(f"{sp.upper()} {lean} leans", gg)
    record = pd.DataFrame(rec_rows) if rec_rows else pd.DataFrame([{"slice": "no settled lines yet"}])
    ws = wb.create_sheet("Real-line record")
    ws["A1"] = "Captured PrizePicks opening lines graded against results (edgeline roi). Break-even for a 4-pick power is 56.2% per leg; the 30% ROI target is 60%. Green = the 95% interval's lower bound clears the bar."
    ws["A1"].font = Font(bold=True, size=12, color=NAVY)
    write_df(ws, record, start_row=3, pct_cols=("hit rate", "CI low", "CI high", "4-pick ROI", "4-pick ROI CI low"))
    if rec_rows:
        n = len(record)
        ws.conditional_formatting.add(f"J4:J{3 + n}", CellIsRule(operator="equal", formula=['"YES"'], fill=PatternFill("solid", fgColor=GREEN)))
        ws.conditional_formatting.add(f"K4:K{3 + n}", CellIsRule(operator="equal", formula=['"YES"'], fill=PatternFill("solid", fgColor=GREEN)))
        ws.conditional_formatting.add(f"E4:E{3 + n}", ColorScaleRule(start_type="num", start_value=0.45, start_color="F8696B", mid_type="num", mid_value=0.5623, mid_color="FFFFFF", end_type="num", end_value=0.7, end_color="63BE7B"))

    # ---------------- Replay ----------------
    rp_rows = []
    for sp in ("dota", "lol", "val", "cs2"):
        p = Path(f"data/backtests/replay_{sp}.csv")
        if not p.exists():
            continue
        d = pd.read_csv(p)
        d["stat"] = d["stat_type"].str.lower().str.replace(r"^map[s]? ?[0-9-]+ ", "", regex=True)
        def addr(name, g):
            n, w = len(g), int(g["win"].sum())
            if n:
                pr, lo, hi = wilson(w, n)
                rp_rows.append({"sport": sp, "slice": name, "n": n, "wins": w, "losses": n - w, "hit rate": pr, "CI low": lo, "CI high": hi, "4-pick ROI": parlay_roi(pr)})
        addr("all leans", d); addr("bettable", d[d["bettable"] == 1])
        for st, g in d.groupby("stat"):
            addr(f"{st} leans", g)
    rp = pd.DataFrame(rp_rows) if rp_rows else pd.DataFrame([{"sport": "no replay yet"}])
    ws = wb.create_sheet("Replay")
    ws["A1"] = "Replay: models trained only on games before 2026-09-25, pointed at that day's real PrizePicks opening lines (out of sample on real lines)."
    ws["A1"].font = Font(bold=True, size=12, color=NAVY)
    write_df(ws, rp, start_row=3, pct_cols=("hit rate", "CI low", "CI high", "4-pick ROI"))
    color_sport(ws, rp, 3)
    if rp_rows:
        ws.conditional_formatting.add(f"F4:F{3 + len(rp)}", ColorScaleRule(start_type="num", start_value=0.45, start_color="F8696B", mid_type="num", mid_value=0.5623, mid_color="FFFFFF", end_type="num", end_value=0.7, end_color="63BE7B"))

    # ---------------- Backtest ----------------
    runs = [("Dota 2", "kills", "dota_kills_bias", "single map", "2 years"), ("Dota 2", "kills", "dota_kills_2map", "two-map sums", "2 years"),
            ("Valorant", "kills", "val_kills_bias", "single map", "7 months"), ("Valorant", "kills", "val_kills_14mo", "single map", "14 months"),
            ("LoL", "kills", "lol_kills_bias", "single map", "2 seasons, all leagues"),
            ("CS2", "kills", "cs2_kills_18mo", "single map", "18 months"), ("CS2", "kills", "cs2_kills_2map", "two-map sums", "18 months"),
            ("CS2", "headshots", "cs2_headshots_18mo", "single map", "18 months"), ("CS2", "headshots", "cs2_headshots_2map", "two-map sums", "18 months"),
            ("COD", "kills", "cod_kills_2seasons", "single map", "2 seasons"),
            ("Valorant", "kills", "val_kills_2map", "two-map sums", "14 months"), ("LoL", "kills", "lol_kills_2map", "two-map sums", "2 seasons")]
    bt_rows = []
    for sport, stat, f, shape, hist in runs:
        p = Path(f"data/backtests/{f}.csv")
        if not p.exists():
            continue
        b = pd.read_csv(p)
        for book in ("booklike", "naive"):
            g = b[b["book"] == book]; sel = g[g["pick"]]
            if not len(sel):
                continue
            n, h = len(sel), int(sel["hit"].sum()); pr, lo, hi = wilson(h, n)
            bt_rows.append({"market": f"{sport} {stat}", "line shape": shape, "history": hist, "setter": "fair book-like model" if book == "booklike" else "naive trailing mean",
                            "games": len(g), "picks (>=60%)": n, "pick share": n / max(len(g), 1), "hit rate": pr, "CI low": lo, "CI high": hi,
                            "4-pick ROI": parlay_roi(pr), "clears 60% (lower bound)": "YES" if lo > 0.60 else "no"})
    bt = pd.DataFrame(bt_rows)
    ws = wb.create_sheet("Backtest")
    ws["A1"] = "Walk-forward backtests (retrain monthly on prior games only, price every player-game, bet the side >= 60%). Lines set at the median-fair half by two synthetic setters; the book-like row is the one to trust. Not realized ROI."
    ws["A1"].font = Font(bold=True, size=12, color=NAVY)
    write_df(ws, bt, start_row=3, pct_cols=("pick share", "hit rate", "CI low", "CI high", "4-pick ROI"))
    if len(bt):
        n = len(bt)
        ws.conditional_formatting.add(f"L4:L{3 + n}", CellIsRule(operator="equal", formula=['"YES"'], fill=PatternFill("solid", fgColor=GREEN)))
        ws.conditional_formatting.add(f"H4:H{3 + n}", ColorScaleRule(start_type="num", start_value=0.55, start_color="F8696B", mid_type="num", mid_value=0.60, mid_color="FFFFFF", end_type="num", end_value=0.70, end_color="63BE7B"))

    # ---------------- Graded lines ----------------
    if not rf.empty:
        gl = rf.copy()
        gl["start (ET)"] = gl["start_time"].map(to_et)
        gl["result"] = gl["win_open"].map({1.0: "WIN", 0.0: "LOSS"}).fillna(gl["result_open"].str.upper())
        gl = gl[["sport", "start (ET)", "player_name", "stat_type", "open_line", "lean", "prob", "ev", "bettable", "actual", "result_open", "result", "model_version"]].sort_values("start (ET)")
        gl["bettable"] = gl["bettable"].map({1: "YES", 0: "no"})
    else:
        gl = pd.DataFrame([{"sport": "no graded lines yet"}])
    ws = wb.create_sheet("Graded lines")
    ws["A1"] = "Every captured line the model leaned on that has settled, graded against the opening line."
    ws["A1"].font = Font(bold=True, size=12, color=NAVY)
    write_df(ws, gl, start_row=3, pct_cols=("prob", "ev"), number_formats={"open_line": "0.0", "actual": "0.0"})
    color_sport(ws, gl, 3)
    if len(gl) and "result" in gl.columns:
        n = len(gl)
        ws.conditional_formatting.add(f"L4:L{3 + n}", CellIsRule(operator="equal", formula=['"WIN"'], fill=PatternFill("solid", fgColor=GREEN)))
        ws.conditional_formatting.add(f"L4:L{3 + n}", CellIsRule(operator="equal", formula=['"LOSS"'], fill=PatternFill("solid", fgColor=RED)))
        ws.conditional_formatting.add(f"L4:L{3 + n}", CellIsRule(operator="equal", formula=['"VOID"'], fill=PatternFill("solid", fgColor=GREY)))
        ws.conditional_formatting.add(f"L4:L{3 + n}", CellIsRule(operator="equal", formula=['"PUSH"'], fill=PatternFill("solid", fgColor=AMBER)))

    # ---------------- Model reference ----------------
    ws = wb.create_sheet("Model")
    ws["A1"] = "Edgeline model reference"; ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    r = 3
    def section(title):
        nonlocal r
        ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=12, color="FFFFFF"); ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=NAVY)
        for c in range(2, 7): ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=NAVY)
        r += 1
    def kv(k, v):
        nonlocal r
        ws.cell(row=r, column=1, value=k).font = BOLD; ws.cell(row=r, column=2, value=v); r += 1
    section("Betting rules (mirrors the original product)")
    kv("Bettable threshold", f"probability >= {DEFAULT_MIN_PROB:.0%} and EV >= {DEFAULT_MIN_EV:.0%} at PrizePicks 4-pick power odds (1.778 per leg)")
    kv("Break-even per leg", "56.2% (4-pick power, 10x); 30% ROI needs 60% per leg (0.60^4 x 10 - 1 = 29.6%)")
    kv("Skips", f"demon/goblin odds; voidable map lines in short series (Riot formats for LoL); lines moved > {DEFAULT_MAX_LINE_MOVE:.0%} against the lean; banned players; started games")
    kv("Market prior", f"{DEFAULT_MARKET_SHRINK:.0%} weight on the book line per component mean")
    kv("Gated markets", ", ".join(f"{s} {t}" for s, t in sorted(UNPROVEN_MARKETS)) or "none (every market clears 60% in backtest as of 2026-09-26)")
    kv("Grading", "opening line vs actual; later-map lines void only when the loaded maps show a decided series")
    r += 1; section("Model")
    kv("Per-map mean", "LightGBM Poisson regression per sport and stat; held-out mean-bias factor over the last 90 days")
    kv("Distribution", "negative binomial (dispersion fit on the recent held-out window); logit calibrator fit on the same window")
    kv("Multi-map and combos", "Gaussian copula over components with rho_self (same player across maps), rho_team, rho_opp; simulated")
    kv("Features", ", ".join(FEATURE_COLUMNS))
    r += 1; section("Payout ladders (net profit per unit by number of correct legs)")
    for book, kinds in LADDERS.items():
        for kind, sizes in kinds.items():
            for size, ladder in sorted(sizes.items()):
                kv(f"{book} {kind} {size}-pick", str(ladder))
    r += 1; section("Data and models")
    for f in sorted(glob.glob("models/artifacts/*_metrics.json")):
        m = json.load(open(f)); name = os.path.basename(f).replace("_metrics.json", "")
        kv(name, f"trained {m.get('n_train', 0):,} rows, valid from {m.get('valid_from')}, MAE {m.get('mae_valid', 0):.3f} vs trailing-mean {m.get('mae_baseline_player_mean10', 0):.3f}, NB r {m.get('dispersion_r', 0):.1f}, mean bias {m.get('mean_bias', 1):.3f}")
    hist = pd.read_sql_query("select sport, min(date) first, max(date) last, count(*) rows, count(distinct game_id) games, count(distinct player_name) players from player_games group by sport", conn)
    for rec in hist.itertuples(index=False):
        kv(f"history {rec.sport}", f"{rec.first[:10]} to {rec.last[:10]}: {rec.rows:,} player-games, {rec.games:,} games, {rec.players:,} players")
    ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 140
    for row in ws.iter_rows(min_row=1, max_row=r, min_col=2, max_col=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size // 1024} KB): board {len(board)}, bettable {len(bet)}, slips {len(slips)}, record {len(record)}, backtest {len(bt)}, graded {len(gl)}")


if __name__ == "__main__":
    main()
